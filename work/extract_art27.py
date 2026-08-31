from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree

source = Path(r"C:\Users\andre\Downloads\2026.08.21-1542_ART. 27 de la RVM 017.docx")
document = Document(source)

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text:
        print(text)

for table_index, table in enumerate(document.tables, start=1):
    print(f"\n[TABLA {table_index}]")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))

with ZipFile(source) as archive:
    root = etree.fromstring(archive.read("word/document.xml"))
    namespaces = {
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    }
    print("\n[ECUACIONES OMML]")
    for index, equation in enumerate(root.xpath(".//m:oMath | .//m:oMathPara", namespaces=namespaces), start=1):
        if equation.getparent() is not None and equation.getparent().tag.endswith("oMathPara") and equation.tag.endswith("oMath"):
            continue
        tokens = equation.xpath(".//m:t/text() | .//w:t/text()", namespaces=namespaces)
        print(index, " ".join(tokens))
